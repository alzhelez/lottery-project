"""Экспорт данных игры в Excel (.xlsx, все таблицы на отдельных листах) и CSV (одна таблица).

CSV пишется в UTF-8 с BOM и разделителем «;» — так его без настроек открывает Excel
с русскими региональными настройками.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from .drawfile import load_game_draws
from .games import Game, results_dir
from .journal import load as load_journal

KINDS = {
    "draws": "Тиражи",
    "journal": "Журнал прогнозов",
    "methods": "Сравнение методов",
    "tests": "Тесты случайности",
    "godmode": "Режим Бога",
}

STATUS = {"pending": "ждёт тиража", "checked": "сверен", "late": "после тиража", "invalid": "не сверить"}


def _join(blocks) -> list[str]:
    return [", ".join(str(v) for v in b) for b in blocks]


def _read(path: Path) -> dict | None:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


def table(game: Game, kind: str) -> tuple[list[str], list[list]] | None:
    """Заголовки и строки таблицы; None — данных для неё нет."""
    if kind == "draws":
        draws = load_game_draws(game)
        if not draws:
            return None
        width = max(len(d.blocks) for d in draws)
        head = ["Тираж", "Дата"] + [f"Поле {i + 1}" for i in range(width)]
        return head, [[d.no, d.date] + _join(d.blocks) + [""] * (width - len(d.blocks)) for d in draws]
    if kind == "journal":
        entries = load_journal(game.slug)["entries"]
        if not entries:
            return None
        head = ["Тираж", "Источник", "Метод", "Прогноз сделан", "Прогноз", "Выпало", "Угадано",
                "Ожидание случайного", "Статус"]
        rows = []
        for e in sorted(entries, key=lambda e: (e["draw_no"], e["made_at"])):
            rows.append([e["draw_no"], e["source_title"], e.get("detail", ""), e["made_at"].replace("T", " "),
                         " | ".join(_join(e["numbers"])), " | ".join(_join(e.get("actual") or [])),
                         e.get("hits", ""), round(e["expected"], 4) if e.get("expected") is not None else "",
                         STATUS.get(e["status"], e["status"])])
        return head, rows
    if kind in ("methods", "tests"):
        a = _read(results_dir(game.slug) / "analysis.json")
        if not a:
            return None
        if kind == "tests":
            tests = a["randomness"].get("tests") or []
            if not tests:
                return None
            head = ["Блок", "Проверка", "Результат", "p", "p с поправкой", "Значимо"]
            return head, [[t["block"], t["title"], t["detail"], round(t["p_value"], 6), round(t["p_adjusted"], 6),
                           "да" if t["significant"] else "нет"] for t in tests]
        ev = a["evaluation"]
        if ev.get("status") != "ok":
            return None
        head = ["Метод", "Контрольный", "Отбор: угадано", "Отбор: к случайному, %", "Контроль: угадано",
                "Контроль: к случайному, %", "p с поправкой", "Прогноз на следующий тираж"]
        rows = []
        for m in sorted(ev["models"], key=lambda m: -m["selection"]["mean_hits"]):
            rows.append([m["title"], "да" if m["control"] else "", round(m["selection"]["mean_hits"], 4),
                         round(m["selection"]["excess_pct"], 2), round(m["holdout"]["mean_hits"], 4),
                         round(m["holdout"]["excess_pct"], 2),
                         round(m["holdout"]["p_adjusted"], 4) if "p_adjusted" in m["holdout"] else "",
                         " | ".join(_join(m["next"]))])
        return head, rows
    if kind == "godmode":
        g = _read(results_dir(game.slug) / "godmode.json")
        if not g:
            return None
        head = ["Тираж", "Дата", "Прогноз", "Выпало", "Угадано", "Модель"]
        return head, [[t["no"], t["date"], " | ".join(_join(t["predicted"])), " | ".join(_join(t["actual"])),
                       t["hits"], t["model"] or "голосование"] for t in g["timeline"]]
    raise ValueError(f"Неизвестный тип экспорта: {kind}")


def to_csv(game: Game, kind: str, out: Path) -> Path:
    t = table(game, kind)
    if t is None:
        raise ValueError(f"Нет данных: {KINDS[kind].lower()}")
    out.parent.mkdir(parents=True, exist_ok=True)
    with open(out, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";")
        w.writerow(t[0])
        # десятичная запятая: иначе русский Excel прочитает 1.62 как дату или текст
        w.writerows([[str(v).replace(".", ",") if isinstance(v, float) else v for v in row] for row in t[1]])
    return out


def to_xlsx(game: Game, out: Path, kinds: list[str] | None = None) -> Path:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    wb.remove(wb.active)
    for kind in kinds or list(KINDS):
        t = table(game, kind)
        if t is None:
            continue
        ws = wb.create_sheet(KINDS[kind][:31])
        ws.append(t[0])
        for row in t[1]:
            ws.append(row)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="E8EDF3")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for i, head in enumerate(t[0], start=1):
            sample = [len(str(r[i - 1])) for r in t[1][:500]] or [0]
            ws.column_dimensions[get_column_letter(i)].width = min(60, max(10, len(head) + 2, max(sample) + 2))
    if not wb.sheetnames:
        raise ValueError("Нет данных для экспорта")
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
