"""Экспорт, автозапуск и продолжение режима Бога."""
import csv
import json
import os
import subprocess
import sys
from datetime import datetime, timedelta

import numpy as np
from conftest import ROOT, random_lotto, write_lotto_file

from lottery.drawfile import Draw, load_game_data, write_draws
from lottery.export import table, to_csv, to_xlsx
from lottery.games import draws_path, get_game, results_dir
from lottery.godmode import ExpertPool, load_state, run_godmode, state_info


def run(args, data_dir, timeout=600):
    env = {**os.environ, "LOTTERY_DATA_DIR": str(data_dir)}
    out = subprocess.run([sys.executable, str(ROOT / "python" / "lp.py")] + args, capture_output=True,
                         text=True, encoding="utf-8", env=env, timeout=timeout)
    events = [json.loads(l[2:]) for l in out.stdout.splitlines() if l.startswith("@@")]
    return out.returncode, events


def fresh_draws(n: int, k: int = 5, n_max: int = 36) -> list[Draw]:
    now = datetime.now().replace(second=0, microsecond=0)
    x = random_lotto(n, n_max, k, 11)
    return [Draw(i + 1, (now - timedelta(hours=n - i)).strftime("%d.%m.%Y %H:%M"),
                 (tuple(int(v) + 1 for v in np.nonzero(x[i])[0]),)) for i in range(n)]


def test_export_csv_and_xlsx(data_dir):
    write_lotto_file(draws_path("5x36"), [random_lotto(300, 36, 5, 1)])
    game = get_game("5x36")
    assert table(game, "journal") is None and table(game, "methods") is None
    out = to_csv(game, "draws", data_dir / "d.csv")
    raw = out.read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")                     # BOM — Excel узнаёт UTF-8
    rows = list(csv.reader(open(out, encoding="utf-8-sig"), delimiter=";"))
    assert rows[0] == ["Тираж", "Дата", "Поле 1"] and len(rows) == 301
    rc, _ = run(["analyze", "--game", "5x36"], data_dir)
    assert rc == 0
    from openpyxl import load_workbook
    wb = load_workbook(to_xlsx(game, data_dir / "all.xlsx"))
    assert wb.sheetnames == ["Тиражи", "Сравнение методов", "Тесты случайности"]
    to_csv(game, "methods", data_dir / "m.csv")
    text = (data_dir / "m.csv").read_text(encoding="utf-8-sig")
    assert "1,6" in text and "1.6" not in text                 # десятичная запятая


def test_auto_cycle_predicts_active_games_only(data_dir):
    write_draws(draws_path("5x36"), fresh_draws(300))
    write_lotto_file(draws_path("6x45"), [random_lotto(300, 45, 6, 2)])   # старые даты — устаревшая
    rc, events = run(["auto", "--skip-fetch"], data_dir)
    assert rc == 0
    status = json.loads((data_dir / "auto_status.json").read_text(encoding="utf-8"))
    assert status["status"] == "done"
    assert status["games"]["5x36"]["analyzed"] and status["games"]["6x45"]["skipped"] == "игра устарела"
    journal = json.loads((data_dir / "journal" / "5x36.json").read_text(encoding="utf-8"))
    # тираж №301 по интервалу приходится на «сейчас» — прогноз пишется на ближайший будущий
    assert {e["draw_no"] for e in journal["entries"]} == {302}
    assert "Автозапуск: готово" in (data_dir / "logs" / "auto.log").read_text(encoding="utf-8")
    assert not (data_dir / "auto.lock").exists()


def test_auto_skips_when_already_running(data_dir):
    (data_dir / "auto.lock").write_text(json.dumps({"pid": os.getpid()}), encoding="utf-8")
    rc, events = run(["auto", "--skip-fetch"], data_dir)
    assert rc == 0 and events[-1].get("skipped")


def test_auto_takes_over_stuck_cycle(data_dir):
    """Цикл «идёт» больше 3 часов (завис, например, после сна компьютера): новый запуск не ждёт
    его вечно. Чужой процесс с тем же номером не трогаем."""
    import subprocess
    import sys
    from datetime import datetime, timedelta
    other = subprocess.Popen([sys.executable, "-c", "import time; time.sleep(60)"])
    try:
        (data_dir / "auto.lock").write_text(json.dumps(
            {"pid": other.pid, "started_at": (datetime.now() - timedelta(hours=5)).isoformat()}), encoding="utf-8")
        rc, events = run(["auto", "--skip-fetch", "--skip-scan"], data_dir)
        assert rc == 0 and not events[-1].get("skipped")
        log = (data_dir / "logs" / "auto.log").read_text(encoding="utf-8")
        assert "считаю зависшим, остановить не удалось" in log and "Автозапуск: готово" in log
        assert other.poll() is None                         # это не lp.py — процесс жив
    finally:
        other.kill()


def test_stop_stuck_kills_only_auto_cycle(tmp_path):
    import subprocess
    import sys
    import time

    import lp
    fake = tmp_path / "lp.py"
    fake.write_text("import time\ntime.sleep(60)\n", encoding="utf-8")
    proc = subprocess.Popen([sys.executable, str(fake), "auto"])
    try:
        time.sleep(1)
        assert lp.stop_stuck(proc.pid)
        proc.wait(timeout=20)
    finally:
        proc.kill()


def test_godmode_resume_continues_saved_state(data_dir, monkeypatch):
    import lottery.godmode as gm
    monkeypatch.setattr(gm, "MIN_BUDGET_S", 4.0)                 # короткий бюджет…
    monkeypatch.setattr(gm, "STAGNATION_BATCHES", 10 ** 6)       # …и перебор не исчерпывается
    write_lotto_file(draws_path("gm"), [random_lotto(500, 20, 4, seed=5)])
    data = load_game_data(get_game("gm"))
    state = results_dir("gm") / "godmode_state.npz"
    first = run_godmode(data, 0.01, state_path=state)
    assert not first["exhausted"]
    info = state_info(state, data)
    assert info["resumable"] and info["experts"] == first["experts"]
    pool, meta = load_state(state, data)
    assert isinstance(pool, ExpertPool) and len(pool.models) == first["experts"]
    # восстановленный пул даёт те же прогнозы и итог
    resumed = run_godmode(data, 0.01, state_path=state, resume=True)
    assert resumed["resumed"] and resumed["experts"] >= first["experts"]
    assert resumed["elapsed_seconds"] >= info["elapsed"]          # время суммируется с прошлым запуском


def test_resume_refused_after_new_draws(data_dir, monkeypatch):
    import lottery.godmode as gm
    monkeypatch.setattr(gm, "MIN_BUDGET_S", 4.0)
    monkeypatch.setattr(gm, "STAGNATION_BATCHES", 10 ** 6)
    x = random_lotto(500, 20, 4, seed=6)
    write_lotto_file(draws_path("gm2"), [x[:450]])
    state = results_dir("gm2") / "godmode_state.npz"
    run_godmode(load_game_data(get_game("gm2")), 0.01, state_path=state)
    write_lotto_file(draws_path("gm2"), [x])
    data = load_game_data(get_game("gm2"))
    assert not state_info(state, data)["resumable"]
    rc, events = run(["godmode", "--game", "gm2", "--budget", "1", "--resume"], data_dir)
    assert rc == 2 and "новые тиражи" in events[-1]["message"]
    assert not (results_dir("gm2") / "godmode.lock").exists()


def test_pool_prune_keeps_controls_and_uses_first_half_only(data_dir):
    write_lotto_file(draws_path("pr"), [random_lotto(600, 20, 4, seed=8)])
    from lottery.models import expert_grid
    pool = ExpertPool(load_game_data(get_game("pr")))
    for m in expert_grid(1):
        pool.add(m)
    controls = sum(m.control for m in pool.models)
    # подменяем вторую половину у одного эксперта — на отбор это влиять не должно
    before = [m.key for m in pool.models]
    pool.hits[5][pool.half():] = 100
    pool.prune(limit=controls + 3)
    assert sum(m.control for m in pool.models) == controls and len(pool.models) == controls + 3
    assert all(k in before for k in (m.key for m in pool.models))
    # эксперт с «идеальной» второй половиной не получил преимущества при отборе
    first_half = {m.key: s for m, s in zip(pool.models, [pool.first_half_score(i) for i in range(len(pool.models))])}
    assert before[5] not in first_half or first_half[before[5]] >= min(
        v for k, v in first_half.items() if not k.startswith("random"))


def test_exhausted_run_drops_state_but_keeps_summary(data_dir):
    write_lotto_file(draws_path("gx"), [random_lotto(400, 20, 4, seed=9)])
    data = load_game_data(get_game("gx"))
    state = results_dir("gx") / "godmode_state.npz"
    res = run_godmode(data, 1, state_path=state)
    assert res["exhausted"]
    assert not state.exists() and state.with_suffix(".json").exists()
    assert np.load  # numpy доступен


def test_state_is_compact(data_dir):
    write_lotto_file(draws_path("gc"), [random_lotto(400, 20, 4, seed=10)])
    from lottery.models import expert_grid
    pool = ExpertPool(load_game_data(get_game("gc")))
    for m in expert_grid(1)[:5]:
        pool.add(m)
    assert all(h.dtype == np.uint8 for h in pool.hits)


def test_cleanup_old_states(data_dir):
    from lottery.godmode import cleanup_states
    folder = results_dir("old")
    folder.mkdir(parents=True)
    (folder / "godmode_state.npz").write_bytes(b"x")
    (folder / "godmode_state.json").write_text("{}")
    os.utime(folder / "godmode_state.npz", (0, 0))
    assert cleanup_states(results_dir()) == 1
    assert not (folder / "godmode_state.npz").exists()


def test_auto_reuses_heavy_checks_when_archive_barely_changed(data_dir, monkeypatch):
    """Автозапуск каждый час: тесты случайности и проверки берутся из прошлого отчёта, пока
    архив вырос меньше чем на 2%; прогноз и проверка методов — всегда заново."""
    import lp
    from lottery import stats
    draws = fresh_draws(400)
    write_draws(draws_path("5x36"), draws[:-1])
    first = lp.analyze_game("5x36", full=False)
    assert first["heavy"]["reused"] is False
    calls = []
    real = stats.game_randomness
    monkeypatch.setattr(stats, "game_randomness", lambda d: calls.append(d.T) or real(d))
    write_draws(draws_path("5x36"), draws)                     # +1 тираж
    second = lp.analyze_game("5x36", full=False)
    assert second["heavy"]["reused"] is True and not calls
    assert second["randomness"] == first["randomness"] and second["last_draw_no"] == first["last_draw_no"] + 1
    assert second["evaluation"]["draws"] == first["evaluation"]["draws"] + 1
    third = lp.analyze_game("5x36", full=True)                 # «Пересчитать анализ» — всё заново
    assert third["heavy"]["reused"] is False and calls


def test_scan_skips_barely_changed_games(data_dir):
    from lottery.scan import run_scan
    draws = fresh_draws(400)
    write_draws(draws_path("5x36"), draws[:-1])
    assert run_scan(parallel=False)["recomputed"] == 1
    write_draws(draws_path("5x36"), draws)
    assert run_scan(parallel=False)["recomputed"] == 0
    write_draws(draws_path("5x36"), draws + fresh_draws(30)[:0] + [Draw(401 + i, draws[-1].date, draws[i].blocks)
                                                                    for i in range(30)])
    assert run_scan(parallel=False)["recomputed"] == 1


def test_auto_runs_godmode_then_reuses_leader(data_dir):
    """Автозапуск: режим Бога для игры с новыми тиражами (прогноз Б в журнал); в следующем
    цикле, пока не прошло 3 часа, — прогноз лидера прошлого запуска без перебора."""
    draws = fresh_draws(420)
    write_draws(draws_path("5x36"), draws[:-1])
    rc, _ = run(["auto", "--skip-fetch", "--skip-scan", "--godmode-budget", "1"], data_dir)
    assert rc == 0
    gm = json.loads((results_dir("5x36") / "godmode.json").read_text(encoding="utf-8"))
    assert gm["auto"] and gm["chosen"] == "leader_all" and gm["next"]["leader"]["key"]
    entries = json.loads((data_dir / "journal" / "5x36.json").read_text(encoding="utf-8"))["entries"]
    assert {e["source"] for e in entries} == {"analysis_best", "analysis_consensus", "chaos", "godmode", "least", "experiment"}
    write_draws(draws_path("5x36"), draws)
    rc, _ = run(["auto", "--skip-fetch", "--skip-scan", "--godmode-budget", "1"], data_dir)
    assert rc == 0
    assert json.loads((results_dir("5x36") / "godmode.json").read_text(encoding="utf-8"))["generated_at"] == gm["generated_at"]
    entries = json.loads((data_dir / "journal" / "5x36.json").read_text(encoding="utf-8"))["entries"]
    # тиражи раз в час: ближайший будущий тираж тот же — прогноз Б для него заменён прогнозом лидера
    latest = max((e for e in entries if e["source"] == "godmode"), key=lambda e: e["made_at"])
    assert latest["draw_no"] >= gm["next"]["no"] and "лидер за всю историю" in latest["detail"]
    log = (data_dir / "logs" / "auto.log").read_text(encoding="utf-8")
    assert "режим Бога за" in log


def test_auto_skips_games_with_few_draws(data_dir):
    write_draws(draws_path("5x36"), fresh_draws(150))
    rc, _ = run(["auto", "--skip-fetch", "--skip-scan", "--godmode-budget", "0"], data_dir)
    status = json.loads((data_dir / "auto_status.json").read_text(encoding="utf-8"))
    assert rc == 0 and status["games"]["5x36"]["skipped"].startswith("мало тиражей")
    assert not (data_dir / "journal" / "5x36.json").exists()


def test_recommend_uses_expected_return_and_proven_source(data_dir):
    import lp
    from lottery.fsutil import write_text
    write_draws(draws_path("5x36"), fresh_draws(400))
    results_dir().mkdir(parents=True, exist_ok=True)
    write_text(results_dir() / "nextdraw.json", json.dumps({"games": {
        "5x36": {"return": 0.71, "usual_return": 0.5, "p_win": 0.1}, "6x45": {"return": 0.4}}}))
    preds = {"5x36": {"title": "5 из 36", "next_no": 401, "numbers": [[1, 2, 3, 4, 5]],
                      "consensus": [[7, 13, 22, 29, 34]]}}
    recs = lp.recommend(preds)
    assert len(recs) == 1 and recs[0]["return"] == 0.71 and recs[0]["source"] == "analysis_consensus"
    assert recs[0]["numbers"] == [[7, 13, 22, 29, 34]] and recs[0]["proven"] is None


def test_auto_godmode_targets_future_draw_at_record_time(data_dir):
    """Режим Бога в автозапуске идёт после анализа всех игр: если тираж, на который был анализ,
    уже прошёл, прогноз Б пишется на ближайший будущий тираж (раньше 44% таких прогнозов
    журнал помечал «после тиража»)."""
    import argparse
    import lp
    from lottery import journal
    write_draws(draws_path("5x36"), fresh_draws(420))       # раз в час, последний — час назад
    folder = results_dir("5x36")
    folder.mkdir(parents=True)
    (folder / "godmode.json").write_text(json.dumps({
        "generated_at": datetime.now().isoformat(timespec="seconds"), "version": lp.__version__,
        "next": {"leader": {"key": "hot_all", "title": "Горячие за всю историю", "params": {}}}}), encoding="utf-8")
    preds = {"5x36": {"title": "5 из 36", "next_no": 421}}  # №421 приходится на «сейчас»
    args = argparse.Namespace(godmode_budget=5, godmode_every=3)
    lp.auto_godmode(args, preds, lambda *a: None, lambda m: None)
    entry = next(e for e in journal.load("5x36")["entries"] if e["source"] == "godmode")
    target = journal.target_draw(load_game_data(get_game("5x36")))["no"]
    assert entry["draw_no"] == target > 421 and preds["5x36"]["godmode_no"] == target
